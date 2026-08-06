"""Append the full operating models (top 8) to the workbook. Run after build_workbook.py."""
import json
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

WB = '/home/user/CI_Test/research/opportunities.xlsx'
fb = json.load(open('/home/user/CI_Test/research/fullbuild.json'))
sourced, models, stress, built = fb['sourced'], fb['models'], fb['stress'], fb['built']
names = {x['id']: x for x in json.load(open('/home/user/CI_Test/research/dataset.json'))['selection']['final_25']}

FONT = 'Arial'
NAVY = '1F3864'
MONEY = '$#,##0;($#,##0);-'
PCT = '0.0%'
HDR_FILL = PatternFill('solid', fgColor=NAVY)
YELLOW = PatternFill('solid', fgColor='FFFF00')
BAND = PatternFill('solid', fgColor='F2F5FA')
INPUT_FONT = Font(name=FONT, size=10, color='0000FF')
THIN = Side(style='thin', color='D0D7E5')
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
STAND_FILL = {True: PatternFill('solid', fgColor='C6EFCE'), False: PatternFill('solid', fgColor='FFC7CE')}
STAND_FONT = {True: Font(name=FONT, size=10, color='006100'), False: Font(name=FONT, size=10, color='9C0006')}

wb = load_workbook(WB)
for nm in ['Model Summary', 'Model Detail', 'Sourcing', 'Corrections', '30-Day Tests']:
    if nm in wb.sheetnames:
        del wb[nm]


def hdr(ws, row, labels, widths_map=None):
    for i, h in enumerate(labels, start=1):
        c = ws.cell(row=row, column=i, value=h)
        c.font = Font(name=FONT, size=10, bold=True, color='FFFFFF')
        c.fill = HDR_FILL
        c.alignment = Alignment(horizontal='left', vertical='center', wrap_text=True)
        c.border = BORDER
    ws.row_dimensions[row].height = 34
    if widths_map:
        for col, w in widths_map.items():
            ws.column_dimensions[col].width = w


def finish(ws, r0, r1, ncols, wrap=()):
    for r in range(r0, r1 + 1):
        for c in range(1, ncols + 1):
            cell = ws.cell(row=r, column=c)
            if cell.font.name != FONT or cell.font.color is None:
                if cell.font.color is None:
                    cell.font = Font(name=FONT, size=10)
            cell.border = BORDER
            cell.alignment = Alignment(vertical='top', wrap_text=(get_column_letter(c) in wrap))
        if (r - r0) % 2 == 1:
            for c in range(1, ncols + 1):
                cc = ws.cell(row=r, column=c)
                if cc.fill.fgColor.rgb in (None, '00000000'):
                    cc.fill = BAND


order = [b['id'] for b in built]

# =====================================================================
# Model Detail — one row per revenue line
# =====================================================================
ws = wb.create_sheet('Model Detail')
ws.sheet_view.showGridLines = False
ws['A1'] = ('Every revenue line modelled separately. Blue cells are inputs. Revenue = units x price; gross profit '
            '= revenue x margin. Model Summary rolls these up with SUMIF, so adding a line here flows straight through.')
ws['A1'].font = Font(name=FONT, size=10, bold=True, color=NAVY)
ws['A1'].alignment = Alignment(wrap_text=True, vertical='top')
ws.merge_cells('A1:V1')
ws.row_dimensions[1].height = 30

D = ['ID', 'Rank', 'Opportunity', 'Revenue line', 'One unit is', 'Price ($)',
     'Units Y1', 'Units Y2', 'Units Y3', 'Units Y4', 'Units Y5',
     'Rev Y1', 'Rev Y2', 'Rev Y3', 'Rev Y4', 'Rev Y5', 'GM %',
     'GP Y1', 'GP Y2', 'GP Y3', 'GP Y4', 'GP Y5']
hdr(ws, 3, D, {'A': 34, 'B': 6, 'C': 34, 'D': 34, 'E': 30, 'F': 12,
               'G': 9, 'H': 9, 'I': 9, 'J': 9, 'K': 9,
               'L': 13, 'M': 13, 'N': 13, 'O': 13, 'P': 13, 'Q': 9,
               'R': 13, 'S': 13, 'T': 13, 'U': 13, 'V': 13})
row = 4
DET_FIRST = row
for b in built:
    oid = b['id']
    for ln in models[oid]['revenue_lines']:
        ws.cell(row=row, column=1, value=oid)
        ws.cell(row=row, column=2, value=b['rank'])
        ws.cell(row=row, column=3, value=names[oid]['name'])
        ws.cell(row=row, column=4, value=ln['name'])
        ws.cell(row=row, column=5, value=ln['unit'])
        c = ws.cell(row=row, column=6, value=ln['price_usd'])
        c.font, c.fill, c.number_format = INPUT_FONT, YELLOW, MONEY
        for i, k in enumerate(['units_y1', 'units_y2', 'units_y3', 'units_y4', 'units_y5']):
            c = ws.cell(row=row, column=7 + i, value=ln[k])
            c.font, c.fill, c.number_format = INPUT_FONT, YELLOW, '#,##0'
        for i in range(5):
            u = get_column_letter(7 + i)
            c = ws.cell(row=row, column=12 + i, value=f'={u}{row}*$F{row}')
            c.number_format = MONEY
        c = ws.cell(row=row, column=17, value=ln['gross_margin_pct'])
        c.font, c.fill, c.number_format = INPUT_FONT, YELLOW, PCT
        for i in range(5):
            rv = get_column_letter(12 + i)
            c = ws.cell(row=row, column=18 + i, value=f'={rv}{row}*$Q{row}')
            c.number_format = MONEY
        row += 1
DET_LAST = row - 1
finish(ws, DET_FIRST, DET_LAST, len(D), wrap=('C', 'D', 'E'))
ws.freeze_panes = 'D4'
ws.auto_filter.ref = f'A3:{get_column_letter(len(D))}{DET_LAST}'

# =====================================================================
# Model Summary — full P&L per opportunity
# =====================================================================
ws = wb.create_sheet('Model Summary')
ws.sheet_view.showGridLines = False
ws['A1'] = ('Full operating model per opportunity, built on LIVE-SOURCED drivers (August 2026). Revenue and gross '
            'profit roll up from Model Detail. Payroll = headcount x loaded salary. EBITDA = gross profit - payroll '
            '- opex. Blue cells are inputs. "Stands?" is the sourcing analyst\'s ruling AFTER live competitor '
            'checks - read the Sourcing and Corrections sheets before trusting any row marked No.')
ws['A1'].font = Font(name=FONT, size=10, bold=True, color='9C0006')
ws['A1'].alignment = Alignment(wrap_text=True, vertical='top')
ws.merge_cells('A1:AR1')
ws.row_dimensions[1].height = 46

S_COLS = (['Rank', 'Opportunity', 'ID', 'Stands?', 'Crowding', 'Confidence'] +
          [f'Revenue Y{i}' for i in range(1, 6)] +
          [f'Gross profit Y{i}' for i in range(1, 6)] +
          [f'Headcount Y{i}' for i in range(1, 6)] +
          ['Loaded salary'] + [f'Payroll Y{i}' for i in range(1, 6)] +
          ['Opex Y1 (input)', 'Opex growth'] + [f'Opex Y{i}' for i in range(2, 6)] +
          [f'EBITDA Y{i}' for i in range(1, 6)] +
          ['Capital required', 'CAC', 'Churn', 'Bear Y5', 'Exit thesis'])
idx = {n: i + 1 for i, n in enumerate(S_COLS)}
hdr(ws, 3, S_COLS)
for n, w in {'Rank': 6, 'Opportunity': 40, 'ID': 32, 'Stands?': 9, 'Crowding': 11, 'Confidence': 11,
             'Loaded salary': 13, 'Opex Y1 (input)': 14, 'Opex growth': 11,
             'Capital required': 15, 'CAC': 11, 'Churn': 9, 'Bear Y5': 14, 'Exit thesis': 70}.items():
    ws.column_dimensions[get_column_letter(idx[n])].width = w
for n in S_COLS:
    if n.startswith(('Revenue', 'Gross profit', 'Payroll', 'Opex Y', 'EBITDA')):
        ws.column_dimensions[get_column_letter(idx[n])].width = 14
    if n.startswith('Headcount'):
        ws.column_dimensions[get_column_letter(idx[n])].width = 11

row = 4
SUM_FIRST = row
for b in built:
    oid = b['id']
    m, sc, st = models[oid], sourced[oid], stress.get(oid, {})
    L = lambda n: get_column_letter(idx[n])
    ws.cell(row=row, column=idx['Rank'], value=b['rank'])
    ws.cell(row=row, column=idx['Opportunity'], value=names[oid]['name'])
    ws.cell(row=row, column=idx['ID'], value=oid)
    stands = bool(sc['still_stands'])
    c = ws.cell(row=row, column=idx['Stands?'], value='Yes' if stands else 'No')
    c.fill, c.font = STAND_FILL[stands], STAND_FONT[stands]
    ws.cell(row=row, column=idx['Crowding'], value=sc['crowding_verdict'])
    cc = ws.cell(row=row, column=idx['Confidence'], value=m['confidence'])
    cc.font = Font(name=FONT, size=10, bold=True, color='9C0006' if m['confidence'] <= 3 else '9C6500')
    # revenue / gross profit roll-ups
    for i in range(1, 6):
        rv = ws.cell(row=row, column=idx[f'Revenue Y{i}'],
                     value=f"=SUMIF('Model Detail'!$A:$A,$C{row},'Model Detail'!{get_column_letter(11+i)}:{get_column_letter(11+i)})")
        rv.number_format = MONEY
        gp = ws.cell(row=row, column=idx[f'Gross profit Y{i}'],
                     value=f"=SUMIF('Model Detail'!$A:$A,$C{row},'Model Detail'!{get_column_letter(17+i)}:{get_column_letter(17+i)})")
        gp.number_format = MONEY
    for i in range(1, 6):
        c = ws.cell(row=row, column=idx[f'Headcount Y{i}'], value=m[f'headcount_y{i}'])
        c.font, c.fill = INPUT_FONT, YELLOW
    c = ws.cell(row=row, column=idx['Loaded salary'], value=m['avg_loaded_salary_usd'])
    c.font, c.fill, c.number_format = INPUT_FONT, YELLOW, MONEY
    for i in range(1, 6):
        p = ws.cell(row=row, column=idx[f'Payroll Y{i}'],
                    value=f"={L(f'Headcount Y{i}')}{row}*${L('Loaded salary')}{row}")
        p.number_format = MONEY
    c = ws.cell(row=row, column=idx['Opex Y1 (input)'], value=m['non_payroll_opex_y1_usd'])
    c.font, c.fill, c.number_format = INPUT_FONT, YELLOW, MONEY
    c = ws.cell(row=row, column=idx['Opex growth'], value=m['non_payroll_opex_growth_pct'])
    c.font, c.fill, c.number_format = INPUT_FONT, YELLOW, PCT
    for i in range(2, 6):
        prev = L('Opex Y1 (input)') if i == 2 else L(f'Opex Y{i-1}')
        o = ws.cell(row=row, column=idx[f'Opex Y{i}'],
                    value=f"={prev}{row}*(1+${L('Opex growth')}{row})")
        o.number_format = MONEY
    for i in range(1, 6):
        opex = L('Opex Y1 (input)') if i == 1 else L(f'Opex Y{i}')
        e = ws.cell(row=row, column=idx[f'EBITDA Y{i}'],
                    value=f"={L(f'Gross profit Y{i}')}{row}-{L(f'Payroll Y{i}')}{row}-{opex}{row}")
        e.number_format = MONEY
        e.font = Font(name=FONT, size=10, bold=True)
    c = ws.cell(row=row, column=idx['Capital required'], value=m['capital_required_usd'])
    c.number_format, c.font, c.fill = MONEY, INPUT_FONT, YELLOW
    c = ws.cell(row=row, column=idx['CAC'], value=m['cac_usd'])
    c.number_format = MONEY
    c = ws.cell(row=row, column=idx['Churn'], value=m['annual_churn_pct'])
    c.number_format = PCT
    c = ws.cell(row=row, column=idx['Bear Y5'], value=st.get('bear_y5_revenue_usd'))
    c.number_format = MONEY
    ws.cell(row=row, column=idx['Exit thesis'], value=m['exit_thesis'])
    ws.row_dimensions[row].height = 54
    row += 1
SUM_LAST = row - 1
ws.cell(row=row, column=idx['Opportunity'], value='Total (8 opportunities)').font = Font(name=FONT, size=10, bold=True)
for n in S_COLS:
    if n.startswith(('Revenue', 'Gross profit', 'EBITDA')) or n in ('Capital required', 'Bear Y5'):
        Lc = get_column_letter(idx[n])
        t = ws.cell(row=row, column=idx[n], value=f'=SUM({Lc}{SUM_FIRST}:{Lc}{SUM_LAST})')
        t.number_format, t.font = MONEY, Font(name=FONT, size=10, bold=True)
finish(ws, SUM_FIRST, row, len(S_COLS), wrap=('B',))
ws.freeze_panes = 'D4'
ws.auto_filter.ref = f'A3:{get_column_letter(len(S_COLS))}{SUM_LAST}'

# =====================================================================
# Sourcing
# =====================================================================
ws = wb.create_sheet('Sourcing')
ws.sheet_view.showGridLines = False
ws['A1'] = ('What live searching (August 2026) established. This is the competitor check the original study never '
            'ran. Only 2 of 8 opportunities survived it intact.')
ws['A1'].font = Font(name=FONT, size=10, bold=True, color='9C0006')
ws['A1'].alignment = Alignment(wrap_text=True, vertical='top')
ws.merge_cells('A1:L1')
ws.row_dimensions[1].height = 30
C = ['Rank', 'Opportunity', 'Trigger verified?', 'What the trigger actually does', 'Crowding',
     'Named competitors found', 'Crowding reasoning', 'Buyer universe (sourced)', 'How it was derived',
     'Pricing comparables found', 'Still stands?', 'Ruling', 'Searches']
hdr(ws, 3, C, {'A': 6, 'B': 36, 'C': 15, 'D': 95, 'E': 11, 'F': 95, 'G': 80, 'H': 15,
               'I': 80, 'J': 80, 'K': 11, 'L': 95, 'M': 9})
row = 4
for b in built:
    sc = sourced[b['id']]
    ws.cell(row=row, column=1, value=b['rank'])
    ws.cell(row=row, column=2, value=names[b['id']]['name'])
    ws.cell(row=row, column=3, value=sc['trigger_verified'])
    ws.cell(row=row, column=4, value=sc['trigger_findings'])
    ws.cell(row=row, column=5, value=sc['crowding_verdict'])
    ws.cell(row=row, column=6, value='\n\n'.join(sc['competitors']))
    ws.cell(row=row, column=7, value=sc['crowding_reasoning'])
    u = ws.cell(row=row, column=8, value=sc['buyer_universe_sourced'])
    u.number_format = '#,##0'
    ws.cell(row=row, column=9, value=sc['buyer_universe_method'])
    ws.cell(row=row, column=10, value=sc['pricing_comparables'])
    st = bool(sc['still_stands'])
    c = ws.cell(row=row, column=11, value='Yes' if st else 'No')
    c.fill, c.font = STAND_FILL[st], STAND_FONT[st]
    ws.cell(row=row, column=12, value=sc['still_stands_reasoning'])
    ws.cell(row=row, column=13, value=sc['searches_run'])
    ws.row_dimensions[row].height = 150
    row += 1
finish(ws, 4, row - 1, len(C), wrap=('B', 'D', 'F', 'G', 'I', 'J', 'L'))
ws.freeze_panes = 'C4'

# =====================================================================
# Corrections
# =====================================================================
ws = wb.create_sheet('Corrections')
ws.sheet_view.showGridLines = False
n_corr = sum(len(sourced[b['id']]['corrections']) for b in built)
ws['A1'] = (f'{n_corr} factual corrections that live sourcing made to the original research. The original study '
            'ran with no live search for these checks, so this is where it was wrong. Read this before acting on '
            'anything in the earlier sheets.')
ws['A1'].font = Font(name=FONT, size=11, bold=True, color='9C0006')
ws['A1'].alignment = Alignment(wrap_text=True, vertical='top')
ws.merge_cells('A1:D1')
ws.row_dimensions[1].height = 34
hdr(ws, 3, ['Rank', 'Opportunity', '#', 'Correction'], {'A': 6, 'B': 36, 'C': 5, 'D': 150})
row = 4
for b in built:
    for i, corr in enumerate(sourced[b['id']]['corrections'], start=1):
        ws.cell(row=row, column=1, value=b['rank'])
        ws.cell(row=row, column=2, value=names[b['id']]['name'])
        ws.cell(row=row, column=3, value=i)
        ws.cell(row=row, column=4, value=corr)
        ws.row_dimensions[row].height = 78
        row += 1
finish(ws, 4, row - 1, 4, wrap=('B', 'D'))
ws.freeze_panes = 'C4'
ws.auto_filter.ref = f'A3:D{row-1}'

# =====================================================================
# 30-Day Tests
# =====================================================================
ws = wb.create_sheet('30-Day Tests')
ws.sheet_view.showGridLines = False
ws['A1'] = ('The cheapest experiment that would falsify each thesis inside 30 days. This is the most actionable '
            'sheet in the workbook - it tells you what to do on Monday instead of what to believe.')
ws['A1'].font = Font(name=FONT, size=11, bold=True, color=NAVY)
ws['A1'].alignment = Alignment(wrap_text=True, vertical='top')
ws.merge_cells('A1:G1')
ws.row_dimensions[1].height = 30
hdr(ws, 3, ['Rank', 'Opportunity', 'Base case defensible?', 'Fatal flaw', 'Weakest number',
            'The 30-day falsification test', 'Bear Y5'],
    {'A': 6, 'B': 34, 'C': 13, 'D': 90, 'E': 85, 'F': 110, 'G': 14})
row = 4
for b in built:
    st = stress.get(b['id'], {})
    ws.cell(row=row, column=1, value=b['rank'])
    ws.cell(row=row, column=2, value=names[b['id']]['name'])
    ok = bool(st.get('base_is_defensible'))
    c = ws.cell(row=row, column=3, value='Yes' if ok else 'No')
    c.fill, c.font = STAND_FILL[ok], STAND_FONT[ok]
    ws.cell(row=row, column=4, value=st.get('fatal_flaw'))
    ws.cell(row=row, column=5, value=st.get('weakest_number'))
    ws.cell(row=row, column=6, value=st.get('what_to_test_first'))
    bc = ws.cell(row=row, column=7, value=st.get('bear_y5_revenue_usd'))
    bc.number_format = MONEY
    ws.row_dimensions[row].height = 150
    row += 1
finish(ws, 4, row - 1, 7, wrap=('B', 'D', 'E', 'F'))
ws.freeze_panes = 'C4'

wb.calculation.fullCalcOnLoad = True
wb.save(WB)
print('sheets:', wb.sheetnames)
print('revenue lines:', DET_LAST - DET_FIRST + 1, '| corrections:', n_corr)
